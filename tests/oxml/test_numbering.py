# encoding: utf-8

"""
Test suite for the list labels computed in ``docx.oxml.numbering``.
"""

from __future__ import absolute_import, print_function, unicode_literals

import pytest

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


class DescribeCT_Numbering(object):

    def it_computes_multi_level_labels(self, labels_fixture):
        document, expected_labels = labels_fixture
        assert [p.number for p in document.paragraphs] == expected_labels

    def it_counts_list_items_as_word_does(self, counting_fixture):
        document, expected_labels = counting_fixture
        assert [p.number for p in document.paragraphs] == expected_labels

    # fixtures -------------------------------------------------------

    @pytest.fixture(params=[
        # Word shows %1 as nothing when level 1 has numFmt="none"
        ('none', '%12030.%2', ['2030.1\t', '2030.2\t']),
        ('decimal', '%1.%2', ['2030.1\t', '2030.2\t']),
    ])
    def labels_fixture(self, request):
        lvl0_numFmt, lvl1_text, expected_labels = request.param
        document = Document()
        numbering = document.part.numbering_part.element
        numbering.append(parse_xml(
            '<w:abstractNum %s w:abstractNumId="90">'
            '  <w:lvl w:ilvl="0"><w:start w:val="2030"/>'
            '    <w:numFmt w:val="%s"/><w:lvlText w:val="2030"/></w:lvl>'
            '  <w:lvl w:ilvl="1"><w:start w:val="1"/>'
            '    <w:numFmt w:val="decimal"/><w:lvlText w:val="%s"/></w:lvl>'
            '</w:abstractNum>' % (nsdecls('w'), lvl0_numFmt, lvl1_text)
        ))
        numbering.append(parse_xml(
            '<w:num %s w:numId="90"><w:abstractNumId w:val="90"/></w:num>'
            % nsdecls('w')
        ))
        for text in ('First.', 'Second.'):
            paragraph = document.add_paragraph(text)
            paragraph._p.get_or_add_pPr().append(parse_xml(
                '<w:numPr %s><w:ilvl w:val="1"/><w:numId w:val="90"/></w:numPr>'
                % nsdecls('w')
            ))
        return document, expected_labels

    # Aspose and LibreOffice show these labels. LibreOffice does not use
    # w:lvlRestart; Aspose does, as the standard says.
    @pytest.fixture(params=[
        # w:num elements on one w:abstractNum share its counters
        ({1: {}, 2: {}}, [(1, 0), (1, 0), (2, 0), (1, 0)],
         ['1.', '2.', '3.', '4.']),
        # a startOverride restarts the shared counter at the first use of
        # its w:num, and only there
        ({1: {}, 2: {0: 1}}, [(1, 0), (1, 0), (2, 0), (1, 0), (2, 0), (1, 0)],
         ['1.', '2.', '1.', '2.', '3.', '4.']),
        ({1: {}, 2: {0: 5}}, [(1, 0), (2, 0), (1, 0)],
         ['1.', '5.', '6.']),
        ({1: {}, 2: {0: 0}}, [(1, 0), (1, 0), (2, 0), (1, 0)],
         ['1.', '2.', '0.', '1.']),
        # the restart waits for the first use at its own level
        ({1: {}, 2: {0: 1}}, [(1, 0), (1, 0), (2, 1), (2, 0), (1, 0)],
         ['1.', '2.', '(a)', '1.', '2.']),
        # an item of any w:num of the list restarts the levels below it
        ({1: {}, 2: {}}, [(1, 1), (1, 1), (2, 0), (1, 1)],
         ['(a)', '(b)', '2.', '(a)']),
        # an item starts the levels above it
        ({1: {}}, [(1, 1), (1, 1), (1, 0), (1, 1)],
         ['(a)', '(b)', '2.', '(a)']),
        ({1: {}}, [(1, 2), (1, 0)],
         ['(i)', '2.']),
        # lvlRestart 0: level 1 never restarts
        ({3: {}}, [(3, 1), (3, 1), (3, 0), (3, 1)],
         ['(a)', '(b)', '2.', '(c)']),
        # lvlRestart 1: level 2 restarts after level 0 only
        ({3: {}}, [(3, 0), (3, 2), (3, 2), (3, 1), (3, 2), (3, 0), (3, 2)],
         ['1.', '(i)', '(ii)', '(b)', '(iii)', '2.', '(i)']),
        # w:num elements on two w:abstractNum elements do not share
        ({1: {}, 4: {}}, [(1, 0), (1, 0), (4, 0), (1, 0)],
         ['1.', '2.', '1.', '3.']),
    ])
    def counting_fixture(self, request):
        nums, paragraphs, expected_labels = request.param
        document = Document()
        numbering = document.part.numbering_part.element
        for abstractNum_id, restarts in ((91, ('', '')), (92, ('', '')),
                                         (93, ('0', '1'))):
            numbering.append(parse_xml(
                '<w:abstractNum %s w:abstractNumId="%d">'
                '  <w:lvl w:ilvl="0"><w:start w:val="1"/>'
                '    <w:numFmt w:val="decimal"/><w:lvlText w:val="%%1."/></w:lvl>'
                '  <w:lvl w:ilvl="1"><w:start w:val="1"/>%s'
                '    <w:numFmt w:val="lowerLetter"/><w:lvlText w:val="(%%2)"/></w:lvl>'
                '  <w:lvl w:ilvl="2"><w:start w:val="1"/>%s'
                '    <w:numFmt w:val="lowerRoman"/><w:lvlText w:val="(%%3)"/></w:lvl>'
                '</w:abstractNum>' % ((nsdecls('w'), abstractNum_id) + tuple(
                    '<w:lvlRestart w:val="%s"/>' % r if r else '' for r in restarts))
            ))
        abstractNum_of = {1: 91, 2: 91, 3: 93, 4: 92}
        for numId, overrides in nums.items():
            numbering.append(parse_xml(
                '<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/>%s</w:num>' % (
                    nsdecls('w'), 90 + numId, abstractNum_of[numId], ''.join(
                        '<w:lvlOverride w:ilvl="%d"><w:startOverride w:val="%d"/>'
                        '</w:lvlOverride>' % item for item in overrides.items()))
            ))
        for numId, ilvl in paragraphs:
            paragraph = document.add_paragraph('Item.')
            paragraph._p.get_or_add_pPr().append(parse_xml(
                '<w:numPr %s><w:ilvl w:val="%d"/><w:numId w:val="%d"/></w:numPr>'
                % (nsdecls('w'), ilvl, 90 + numId)
            ))
        return document, [label + '\t' for label in expected_labels]
