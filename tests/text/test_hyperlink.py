# encoding: utf-8

"""
Test suite for the docx.text.hyperlink module.
"""

from __future__ import absolute_import, print_function, unicode_literals

from docx import Document


class DescribeHyperlink(object):

    def it_knows_the_url_of_a_hyperlink_in_the_body(self):
        paragraph = Document().add_paragraph('See ')

        hyperlink = paragraph.add_hyperlink('the site', 'https://example.com/body')

        assert hyperlink.link == 'https://example.com/body'

    def it_keeps_the_url_of_a_hyperlink_in_a_footnote_in_the_footnotes_part(self):
        document = Document()
        footnote = document.add_paragraph('See').add_footnote()
        paragraph = footnote.paragraphs[0]

        hyperlink = paragraph.add_hyperlink('the site', 'https://example.com/note')

        rId = hyperlink._h.relationship_id
        assert paragraph.part is not document.part
        assert paragraph.part.rels[rId].target_ref == 'https://example.com/note'
        assert hyperlink.link == 'https://example.com/note'
