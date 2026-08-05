# pyright: reportPrivateUsage=false

"""Unit test suite for the `docx.endnotes` module."""

from __future__ import annotations

from typing import cast

import pytest

from docx.endnotes import Endnote, Endnotes
from docx.oxml.endnote import CT_Endnotes, CT_FtnEdn
from docx.parts.endnotes import EndnotesPart

from .unitutil.cxml import element, xml
from .unitutil.mock import FixtureRequest, Mock, instance_mock


class DescribeEndnotes:
    """Unit-test suite for `docx.endnotes.Endnotes` objects."""

    @pytest.mark.parametrize(
        ("cxml", "count"),
        [
            ("w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0})", 2),
            (
                "w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0},w:endnote{w:id=1})",
                3,
            ),
        ],
    )
    def it_knows_how_many_endnotes_it_contains(self, cxml: str, count: int, endnotes_part_: Mock):
        endnotes_elm = cast(CT_Endnotes, element(cxml))
        endnotes = Endnotes(endnotes_elm, endnotes_part_)

        assert len(endnotes) == count

    def it_can_get_an_endnote_by_reference_id(self, endnotes_part_: Mock):
        endnotes_elm = cast(
            CT_Endnotes,
            element(
                "w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0},"
                "w:endnote{w:id=1},w:endnote{w:id=2})"
            ),
        )
        endnotes = Endnotes(endnotes_elm, endnotes_part_)

        endnote = endnotes[2]

        assert isinstance(endnote, Endnote)
        assert endnote._e is endnotes_elm.endnote_sequence_lst[3]

    def it_raises_when_no_endnote_has_that_reference_id(self, endnotes_part_: Mock):
        endnotes_elm = cast(
            CT_Endnotes, element("w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0})")
        )
        endnotes = Endnotes(endnotes_elm, endnotes_part_)

        with pytest.raises(IndexError):
            endnotes[10]

    def it_can_append_a_new_endnote_when_its_id_is_not_taken(self, endnotes_part_: Mock):
        endnotes_elm = cast(
            CT_Endnotes,
            element("w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0},w:endnote{w:id=1})"),
        )
        endnotes = Endnotes(endnotes_elm, endnotes_part_)

        endnote = endnotes.add_endnote(2)

        assert isinstance(endnote, Endnote)
        assert endnote.id == 2
        assert endnotes_elm.xml == xml(
            "w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0},w:endnote{w:id=1},w:endnote{w:id=2})"
        )

    def it_shifts_up_existing_endnotes_when_inserting_ahead_of_them(self, endnotes_part_: Mock):
        endnotes_elm = cast(
            CT_Endnotes,
            element(
                "w:endnotes/(w:endnote{w:id=-1},w:endnote{w:id=0},"
                "w:endnote{w:id=1},w:endnote{w:id=2},w:endnote{w:id=3})"
            ),
        )
        endnotes = Endnotes(endnotes_elm, endnotes_part_)

        endnote = endnotes.add_endnote(1)

        assert endnote.id == 1
        ids = [e.id for e in endnotes_elm.endnote_sequence_lst]
        assert ids == [-1, 0, 1, 2, 3, 4]

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def endnotes_part_(self, request: FixtureRequest):
        return instance_mock(request, EndnotesPart)


class DescribeEndnote:
    """Unit-test suite for `docx.endnotes.Endnote`."""

    def it_knows_its_reference_id(self, endnotes_part_: Mock):
        endnote_elm = cast(CT_FtnEdn, element("w:endnote{w:id=42}"))
        endnote = Endnote(endnote_elm, endnotes_part_)

        assert endnote.id == 42

    def it_provides_access_to_the_paragraphs_it_contains(self, endnotes_part_: Mock):
        endnote_elm = cast(
            CT_FtnEdn,
            element('w:endnote{w:id=1}/(w:p/w:r/w:t"First para",w:p/w:r/w:t"Second para")'),
        )
        endnote = Endnote(endnote_elm, endnotes_part_)

        paragraphs = endnote.paragraphs

        assert len(paragraphs) == 2
        assert [p.text for p in paragraphs] == ["First para", "Second para"]

    def it_can_have_a_paragraph_added_to_it(self, endnotes_part_: Mock):
        endnote_elm = cast(CT_FtnEdn, element("w:endnote{w:id=1}"))
        endnote = Endnote(endnote_elm, endnotes_part_)

        paragraph = endnote.add_paragraph("Some endnote content.")

        assert paragraph.text == "Some endnote content."
        assert len(endnote.paragraphs) == 1

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def endnotes_part_(self, request: FixtureRequest):
        return instance_mock(request, EndnotesPart)
