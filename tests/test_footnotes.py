# pyright: reportPrivateUsage=false

"""Unit test suite for the `docx.footnotes` module."""

from __future__ import annotations

from typing import cast

import pytest

from docx.footnotes import Footnote, Footnotes
from docx.oxml.footnote import CT_Footnotes, CT_FtnEnd
from docx.parts.footnotes import FootnotesPart

from .unitutil.cxml import element, xml
from .unitutil.mock import FixtureRequest, Mock, instance_mock


class DescribeFootnotes:
    """Unit-test suite for `docx.footnotes.Footnotes` objects."""

    @pytest.mark.parametrize(
        ("cxml", "count"),
        [
            ("w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0})", 2),
            (
                "w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},w:footnote{w:id=1})",
                3,
            ),
        ],
    )
    def it_knows_how_many_footnotes_it_contains(self, cxml: str, count: int, footnotes_part_: Mock):
        footnotes_elm = cast(CT_Footnotes, element(cxml))
        footnotes = Footnotes(footnotes_elm, footnotes_part_)

        assert len(footnotes) == count

    def it_can_get_a_footnote_by_reference_id(self, footnotes_part_: Mock):
        footnotes_elm = cast(
            CT_Footnotes,
            element(
                "w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},"
                "w:footnote{w:id=1},w:footnote{w:id=2})"
            ),
        )
        footnotes = Footnotes(footnotes_elm, footnotes_part_)

        footnote = footnotes[2]

        assert isinstance(footnote, Footnote)
        assert footnote._f is footnotes_elm.footnote_sequence_lst[3]

    def it_raises_when_no_footnote_has_that_reference_id(self, footnotes_part_: Mock):
        footnotes_elm = cast(
            CT_Footnotes, element("w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0})")
        )
        footnotes = Footnotes(footnotes_elm, footnotes_part_)

        with pytest.raises(IndexError):
            footnotes[10]

    def it_can_append_a_new_footnote_when_its_id_is_not_taken(self, footnotes_part_: Mock):
        footnotes_elm = cast(
            CT_Footnotes,
            element("w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},w:footnote{w:id=1})"),
        )
        footnotes = Footnotes(footnotes_elm, footnotes_part_)

        footnote = footnotes.add_footnote(2)

        assert isinstance(footnote, Footnote)
        assert footnote.id == 2
        assert footnotes_elm.xml == xml(
            "w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},w:footnote{w:id=1},"
            "w:footnote{w:id=2})"
        )

    def it_shifts_up_existing_footnotes_when_inserting_ahead_of_them(self, footnotes_part_: Mock):
        footnotes_elm = cast(
            CT_Footnotes,
            element(
                "w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},"
                "w:footnote{w:id=1},w:footnote{w:id=2},w:footnote{w:id=3})"
            ),
        )
        footnotes = Footnotes(footnotes_elm, footnotes_part_)

        footnote = footnotes.add_footnote(1)

        assert footnote.id == 1
        assert footnotes_elm.xml == xml(
            "w:footnotes/(w:footnote{w:id=-1},w:footnote{w:id=0},w:footnote{w:id=1},"
            "w:footnote{w:id=2},w:footnote{w:id=3},w:footnote{w:id=4})"
        )
        # -- the previously id=1,2,3 footnotes are now id=2,3,4, and the new one lands
        # -- physically ahead of them, leaving the whole sequence in id order --
        ids = [f.id for f in footnotes_elm.footnote_sequence_lst]
        assert ids == [-1, 0, 1, 2, 3, 4]

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def footnotes_part_(self, request: FixtureRequest):
        return instance_mock(request, FootnotesPart)


class DescribeFootnote:
    """Unit-test suite for `docx.footnotes.Footnote`."""

    def it_knows_its_reference_id(self, footnotes_part_: Mock):
        footnote_elm = cast(CT_FtnEnd, element("w:footnote{w:id=42}"))
        footnote = Footnote(footnote_elm, footnotes_part_)

        assert footnote.id == 42

    def it_provides_access_to_the_paragraphs_it_contains(self, footnotes_part_: Mock):
        footnote_elm = cast(
            CT_FtnEnd,
            element('w:footnote{w:id=1}/(w:p/w:r/w:t"First para",w:p/w:r/w:t"Second para")'),
        )
        footnote = Footnote(footnote_elm, footnotes_part_)

        paragraphs = footnote.paragraphs

        assert len(paragraphs) == 2
        assert [p.text for p in paragraphs] == ["First para", "Second para"]

    def it_can_have_a_paragraph_added_to_it(self, footnotes_part_: Mock):
        footnote_elm = cast(CT_FtnEnd, element("w:footnote{w:id=1}"))
        footnote = Footnote(footnote_elm, footnotes_part_)

        paragraph = footnote.add_paragraph("Some footnote content.")

        assert paragraph.text == "Some footnote content."
        assert len(footnote.paragraphs) == 1

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def footnotes_part_(self, request: FixtureRequest):
        return instance_mock(request, FootnotesPart)
